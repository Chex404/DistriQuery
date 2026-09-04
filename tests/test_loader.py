import pytest

from distriquery.loader import load_document


def test_load_txt(tmp_path):
    file_path = tmp_path / "doc.txt"
    file_path.write_text("hello distriquery")

    doc = load_document(str(file_path))

    assert doc.text == "hello distriquery"
    assert doc.source == str(file_path)


def test_load_md(tmp_path):
    file_path = tmp_path / "doc.md"
    file_path.write_text("# heading\n\nsome text")

    doc = load_document(str(file_path))

    assert "heading" in doc.text


def test_load_docx(tmp_path):
    from docx import Document as DocxDocument

    file_path = tmp_path / "doc.docx"
    docx_file = DocxDocument()
    docx_file.add_paragraph("First paragraph about DistriQuery.")
    docx_file.add_paragraph("Second paragraph about Kafka.")
    docx_file.save(str(file_path))

    doc = load_document(str(file_path))

    assert "DistriQuery" in doc.text
    assert "Kafka" in doc.text


def test_missing_file_raises(tmp_path):
    missing = tmp_path / "nope.txt"

    with pytest.raises(FileNotFoundError):
        load_document(str(missing))


def test_unsupported_extension_raises(tmp_path):
    file_path = tmp_path / "doc.exe"
    file_path.write_text("binary-ish content")

    with pytest.raises(ValueError):
        load_document(str(file_path))