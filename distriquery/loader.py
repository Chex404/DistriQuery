"""Document loading.

Turns a file on disk into a plain-text Document object. This is intentionally
the dumbest possible layer: no chunking, no cleaning beyond basic decoding.
Later phases (Kafka workers in Phase D) will call load_document() the same
way this module's tests do — that's the point of keeping it isolated.
"""

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}  # later: .html, .pptx, etc.


@dataclass
class Document:
    source: str  # file path (later: could be a URL, S3 key, etc.)
    text: str


def load_document(path: str) -> Document:
    """Load a single document from disk.

    Raises FileNotFoundError if the path doesn't exist, and ValueError for
    an unsupported extension — both are things a caller (or a Kafka worker
    in Phase D) needs to handle explicitly, so we raise rather than return
    None or an empty string.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"No such file: {path}")

    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{extension}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if extension == ".pdf":
        text = _load_pdf(file_path)
    elif extension == ".docx":
        text = _load_docx(file_path)
    else:
        text = file_path.read_text(encoding="utf-8", errors="replace")

    return Document(source=str(file_path), text=text)


def _load_pdf(file_path: Path) -> str:
    # Imported lazily so .txt/.md-only usage never needs pypdf installed.
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _load_docx(file_path: Path) -> str:
    # Imported lazily so .txt/.md/.pdf-only usage never needs python-docx installed.
    from docx import Document as DocxDocument

    docx_file = DocxDocument(str(file_path))
    paragraphs = [p.text for p in docx_file.paragraphs]
    return "\n\n".join(paragraphs)